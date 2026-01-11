from docx import Document
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# Create a new Document
doc = Document()

# Add a paragraph with an inline math formula
p = doc.add_paragraph('The area of a circle is ')
run = p.add_run()
# Add an inline OMML math formula (A=πr²)
mml = parse_xml('''
<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
    <m:r>
        <m:t>A</m:t>
    </m:r>
    <m:r>
        <m:t>=</m:t>
    </m:r>
    <m:r>
        <m:t>π</m:t>
    </m:r>
    <m:r>
        <m:t>r</m:t>
    </m:r>
    <m:r>
        <m:t>²</m:t>
    </m:r>
</m:oMath>
''')
run._r.append(mml)

# Add a paragraph with a block math formula
doc.add_paragraph('The quadratic formula is:')
p = doc.add_paragraph()
run = p.add_run()
# Add a block OMML math formula (x = (-b ± √(b²-4ac))/2a)
mml = parse_xml('''
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
    <m:oMath>
        <m:r>
            <m:t>x</m:t>
        </m:r>
        <m:r>
            <m:t>=</m:t>
        </m:r>
        <m:r>
            <m:t>(</m:t>
        </m:r>
        <m:r>
            <m:t>-b</m:t>
        </m:r>
        <m:r>
            <m:t> ± </m:t>
        </m:r>
        <m:r>
            <m:t>√</m:t>
        </m:r>
        <m:r>
            <m:t>(</m:t>
        </m:r>
        <m:r>
            <m:t>b²-4ac</m:t>
        </m:r>
        <m:r>
            <m:t>)</m:t>
        </m:r>
        <m:r>
            <m:t>)</m:t>
        </m:r>
        <m:r>
            <m:t>/2a</m:t>
        </m:r>
    </m:oMath>
</m:oMathPara>
''')
run._r.append(mml)

# Save the document
doc.save('test-math.docx')
print('test-math.docx created successfully.')